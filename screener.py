"""Resume screening core: text extraction, rubric scoring via LLM, storage, email.

Scores come from the model (0-3 per parameter + a verbatim CV quote); the
weighted total and decision band are computed here in code, never by the model.
"""
import hashlib
import json
import os
import re
import smtplib
import threading
import time
import uuid
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from email.message import EmailMessage
from pathlib import Path

import requests
from dotenv import load_dotenv

ROOT = Path(__file__).parent
load_dotenv(ROOT / ".env")

# On Vercel each request is a short-lived function: no background threads, only /tmp is writable,
# and Supabase is the single source of truth.
SERVERLESS = bool(os.environ.get("VERCEL"))
DATA = ROOT / "data"
UPLOADS = Path("/tmp/uploads") if SERVERLESS else DATA / "uploads"
DB_FILE = DATA / "candidates.json"
REF = ROOT / "reference"

PARAMS = [
    ("p1", "JD baseline fit", 1),
    ("p2", "Ownership, no safety net", 1),
    ("p3", "Ground-level domain", 2),
    ("p4", "Builds unasked, adopted", 2),
    ("p5", "Closes the loop", 1),
    ("p6", "Outcome evidence", 1),
]
MAX_WEIGHTED = sum(3 * w for _, _, w in PARAMS)  # 24

# Calibration table from rubric section 4 (P2-P6; P1 excluded for past hires).
PAST_HIRES = {
    "lavanya": ("Lavanya", "Exceeds", [3, 3, 3, 3, 3]),
    "rohan_desai": ("Rohan", "Exceeds", [3, 3, 3, 2, 3]),
    "meghna": ("Meghna", "Exceeds", [2, 3, 3, 3, 3]),
    "sunita": ("Sunita", "Exceeds", [3, 3, 3, 2, 2]),
    "aditya_shetty": ("Aditya", "Exceeds", [3, 3, 2, 3, 3]),
    "rahul_bose": ("Rahul", "Meets", [2, 0, 2, 0, 3]),
    "preetham": ("Preetham", "Below", [1, 1, 1, 1, 3]),
    "vikram": ("Vikram", "Meets", [1, 0, 2, 1, 2]),
}


def env(name, default=None):
    return os.environ.get(name) or default


def band_for(total):
    if total >= 17:
        return "SHORTLIST"
    if total >= 12:
        return "BORDERLINE"
    return "NOT MOVING FORWARD"


# ---------- text extraction ----------

def extract_text(path):
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        import pypdf
        text = "\n".join((p.extract_text() or "") for p in pypdf.PdfReader(str(path)).pages)
    elif ext == ".docx":
        import docx
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        text = "\n".join(parts)
    elif ext in (".txt", ".md"):
        text = path.read_text(encoding="utf-8", errors="replace")
    else:
        raise ValueError(f"Unsupported file type {ext} (use PDF, DOCX or TXT)")
    text = text.replace("�", "·")
    if len(text.strip()) < 200:
        raise ValueError("Could not read text from this file (scanned image?). Upload a text PDF or DOCX.")
    return text


def _norm(s):
    return re.sub(r"[^a-z0-9%]+", " ", s.lower()).strip()


def evidence_found(evidence, cv_text):
    """True if the quoted evidence really appears in the CV (tolerant of PDF bullet/space noise)."""
    e, cv = _norm(evidence), _norm(cv_text)
    if not e:
        return False
    if e in cv:
        return True
    words = e.split()
    if len(words) < 4:
        return False
    grams = [" ".join(words[i:i + 4]) for i in range(len(words) - 3)]
    return sum(g in cv for g in grams) / len(grams) >= 0.7


# ---------- prompt ----------

def _past_hire_block():
    out = []
    for f in sorted((REF / "past_hires").glob("*.docx")):
        key = next((k for k in PAST_HIRES if k in f.stem), None)
        if not key:
            continue
        name, rating, s = PAST_HIRES[key]
        scores = ", ".join(f"P{i + 2}={v}" for i, v in enumerate(s))
        out.append(f"<past_hire name=\"{name}\" rating=\"{rating}\" rubric_scores=\"{scores}\">\n"
                   f"{extract_text(f).strip()}\n</past_hire>")
    return "\n\n".join(out)


SYSTEM_PROMPT = None


def system_prompt():
    global SYSTEM_PROMPT
    if SYSTEM_PROMPT is None:
        rubric = (REF / "rubric.txt").read_text(encoding="utf-8")
        SYSTEM_PROMPT = f"""You screen applicants for the Product Manager (PM) and Senior Product Manager (Sr PM) roles at Kargo, a logistics / freight-tech company in Mumbai. You apply the hiring rubric below exactly. It was built from the job descriptions and from the CVs of past hires and how they performed.

<rubric>
{rubric}
</rubric>

Below are the CVs of the 8 past hires the rubric was calibrated on, with the scores the hiring team gave them on P2-P6. Use them as anchors: a 3 should look like the evidence in the Exceeds hires, and the Meets/Below hires show what does not earn points (credentials, big-company scale, touching logistics only through APIs).

<calibration>
{_past_hire_block()}
</calibration>

How to score an applicant:
- Score each parameter P1-P6 from 0 to 3 using the rubric table. For P1 use the role notes for the role the applicant is being assessed for.
- For every parameter, "evidence" must be ONE line copied word for word from the CV that justifies the score. If the CV has no evidence for a parameter, the score is 0 and evidence is an empty string. Do not infer, guess, or give credit for what the person probably did.
- Location is a gate inside P1: if the CV shows they are not in Mumbai and gives no sign of relocating, say "Unclear" and note it; score P1 at most 1 only if they are clearly not open to Mumbai.
- Never let name, gender, age, college brand, or city of origin affect any score. Do not reward credentials, certifications, MBAs, or employer size.
- Be strict and consistent: the same evidence should get the same score for every applicant.
- You do not compute the total or the decision band; the system does that from your scores.

Also write, for the hiring manager Arjun:
- why_ranked: two lines on why this applicant lands where they do.
- brief: an interview brief. Strengths and concerns must be grounded in the CV. probe_questions: 2 questions tied to the weakest-scoring parameters (as the rubric asks), then up to 2 more worth asking; each says which parameter it targets and what a strong answer would contain.
- invite_email: a warm, short interview invitation from the sender to the applicant. Mention one specific thing from their CV that stood out. Include the scheduling instruction you are given.
- decline_email: a respectful, kind decline. Thank them, do not quote scores or the rubric, do not give false hope, wish them well. Under 120 words.
For both emails, "paragraphs" holds only the body: 2 or 3 short paragraphs of 1-3 sentences each. Do NOT include a greeting ("Hi ...") or a sign-off ("Best regards", the sender's name); the system adds those. Plain text, no markdown, no placeholders in square brackets except where information is truly unknown."""
    return SYSTEM_PROMPT


def _param_schema(desc):
    return {
        "type": "object",
        "description": desc,
        "properties": {
            "score": {"type": "integer", "description": "0, 1, 2 or 3"},
            "evidence": {"type": "string", "description": "one line copied verbatim from the CV, or empty if none"},
            "reasoning": {"type": "string", "description": "one sentence mapping the evidence to the rubric level"},
        },
        "required": ["score", "evidence", "reasoning"],
    }


def _email_schema():
    return {"type": "object", "properties": {
        "subject": {"type": "string"},
        "paragraphs": {"type": "array", "items": {"type": "string"},
                       "description": "2-3 short body paragraphs; no greeting, no sign-off"}},
        "required": ["subject", "paragraphs"]}


def compose_email(first_name, paragraphs):
    """Lay an email out the same way every time: greeting, paragraphs separated by blank lines, sign-off."""
    paras = [" ".join(p.split()) for p in paragraphs if p and p.strip()]
    return (f"Hi {first_name},\n\n" + "\n\n".join(paras) +
            f"\n\nBest regards,\n{env('SENDER_NAME', 'Arjun')}\n{env('COMPANY_NAME', 'Kargo')}")


SIGNOFF_RE = re.compile(r"\s*\b((?:Best|Warm|Kind|Warmest)\s+regards|Best wishes|Many thanks|Sincerely|Regards|Thanks|Best)\s*,\s*[^.!?]{0,60}$")


def tidy_email(body, first_name):
    """Re-lay out an email that arrived as one run-on block (older drafts): keeps the wording, fixes the format."""
    text = " ".join((body or "").split())
    m = re.match(r"^(?:Hi|Hello|Dear)\s+[^,]{1,40},\s*", text)
    if m:
        text = text[m.end():]
    s = SIGNOFF_RE.search(text)
    if s:
        text = text[:s.start()]
    sentences = [x for x in re.split(r"(?<=[.!?])\s+(?=[A-Z\"'(])", text.strip()) if x]
    # 1 opening sentence, then pairs of sentences, so a typical email gets 2-4 short paragraphs
    paras = sentences[:1] + [" ".join(sentences[i:i + 2]) for i in range(1, len(sentences), 2)]
    return compose_email(first_name, paras)


SCHEMA = {
    "type": "object",
    "properties": {
        "candidate": {
            "type": "object",
            "properties": {
                "name": {"type": "string"},
                "email": {"type": "string", "description": "empty if not on the CV"},
                "phone": {"type": "string"},
                "location": {"type": "string"},
                "current_title": {"type": "string"},
                "total_years_experience": {"type": "number"},
                "pm_years": {"type": "number", "description": "years in product manager roles only"},
                "role_assessed": {"type": "string", "enum": ["PM", "Sr PM"]},
                "role_reason": {"type": "string"},
                "location_gate": {"type": "string", "enum": ["Mumbai", "Willing to relocate", "Not open", "Unclear"]},
            },
            "required": ["name", "email", "phone", "location", "current_title", "total_years_experience",
                         "pm_years", "role_assessed", "role_reason", "location_gate"],
        },
        "scores": {
            "type": "object",
            "properties": {k: _param_schema(label) for k, label, _ in PARAMS},
            "required": [k for k, _, _ in PARAMS],
        },
        "why_ranked": {"type": "string"},
        "brief": {
            "type": "object",
            "properties": {
                "summary": {"type": "string", "description": "3-4 sentence profile"},
                "strengths": {"type": "array", "items": {"type": "string"}},
                "concerns": {"type": "array", "items": {"type": "string"}},
                "probe_questions": {
                    "type": "array",
                    "items": {
                        "type": "object",
                        "properties": {
                            "question": {"type": "string"},
                            "targets": {"type": "string", "description": "e.g. P3 Ground-level domain"},
                            "listen_for": {"type": "string"},
                        },
                        "required": ["question", "targets", "listen_for"],
                    },
                },
            },
            "required": ["summary", "strengths", "concerns", "probe_questions"],
        },
        "invite_email": _email_schema(),
        "decline_email": _email_schema(),
    },
    "required": ["candidate", "scores", "why_ranked", "brief", "invite_email", "decline_email"],
}


def _for_claude(s):
    s = dict(s)
    if s.get("type") == "object":
        s["additionalProperties"] = False
        s["properties"] = {k: _for_claude(v) for k, v in s["properties"].items()}
    if s.get("type") == "array":
        s["items"] = _for_claude(s["items"])
    return s


def _for_gemini(s):
    s = {k: v for k, v in s.items() if k != "additionalProperties"}
    s["type"] = s["type"].upper()
    if "properties" in s:
        s["properties"] = {k: _for_gemini(v) for k, v in s["properties"].items()}
    if "items" in s:
        s["items"] = _for_gemini(s["items"])
    return s


def user_prompt(cv_text, role_hint):
    sender = env("SENDER_NAME", "Arjun")
    company = env("COMPANY_NAME", "Kargo")
    link = env("SCHEDULING_LINK")
    sched = (f"Ask them to pick a slot at {link}" if link
             else "Ask them to reply with three 45-minute slots that work for them over the next week")
    role_line = (f"Assess this applicant for the {role_hint} role." if role_hint else
                 "Decide from the CV whether to assess them for PM (about 2-4 yrs PM) or Sr PM (about 5-8 yrs PM, "
                 "or they explicitly target Senior PM), then assess for that role.")
    return f"""{role_line}
Emails are signed by {sender}, {company}. Scheduling: {sched}.

<cv>
{cv_text}
</cv>"""


# ---------- model calls ----------

class LLMError(RuntimeError):
    pass


def call_llm(cv_text, role_hint):
    provider = env("LLM_PROVIDER", "claude").lower()
    prompt = user_prompt(cv_text, role_hint)
    if provider == "gemini":
        return _gemini(prompt)
    return _claude(prompt)


_claude_client = None


def _claude(prompt):
    global _claude_client
    import anthropic
    if _claude_client is None:
        _claude_client = anthropic.Anthropic()
    try:
        resp = _claude_client.messages.create(
            model=env("CLAUDE_MODEL", "claude-opus-5"),
            max_tokens=16000,
            system=[{"type": "text", "text": system_prompt(), "cache_control": {"type": "ephemeral"}}],
            thinking={"type": "adaptive"},
            output_config={"effort": env("CLAUDE_EFFORT", "medium"),
                           "format": {"type": "json_schema", "schema": _for_claude(SCHEMA)}},
            messages=[{"role": "user", "content": prompt}],
        )
    except anthropic.AuthenticationError:
        raise LLMError("Anthropic API key missing or invalid - set ANTHROPIC_API_KEY in .env")
    except anthropic.RateLimitError:
        raise LLMError("Rate limited by Anthropic - click Rescore in a minute")
    except anthropic.APIStatusError as e:
        raise LLMError(f"Anthropic API error {e.status_code}: {e.message}")
    except anthropic.APIConnectionError:
        raise LLMError("Could not reach the Anthropic API")
    if resp.stop_reason == "refusal":
        raise LLMError("The model declined to score this CV")
    if resp.stop_reason == "max_tokens":
        raise LLMError("Model output was cut off - click Rescore")
    text = "".join(b.text for b in resp.content if b.type == "text")
    return json.loads(text)


def _gemini(prompt):
    key = env("GEMINI_API_KEY")
    if not key:
        raise LLMError("GEMINI_API_KEY missing in .env")
    model = env("GEMINI_MODEL", "gemini-3.5-flash")
    body = {
        "systemInstruction": {"parts": [{"text": system_prompt()}]},
        "contents": [{"role": "user", "parts": [{"text": prompt}]}],
        "generationConfig": {"temperature": 0.2, "responseMimeType": "application/json",
                             "responseSchema": _for_gemini(SCHEMA)},
    }
    url = f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent"
    for attempt in range(3):  # dropped connections, rate limits and 5xx are usually gone a few seconds later
        try:
            r = requests.post(url, json=body, headers={"x-goog-api-key": key}, timeout=180)
            if r.status_code not in (429, 500, 502, 503, 504) or attempt == 2:
                break
        except requests.RequestException as e:
            if attempt == 2:
                raise LLMError(f"Could not reach Gemini: {e}")
        time.sleep(3 * (attempt + 1))
    try:
        data = r.json()
    except ValueError:
        raise LLMError(f"Gemini error {r.status_code}")
    if r.status_code != 200:
        raise LLMError(f"Gemini error {r.status_code}: {data.get('error', {}).get('message', data)}")
    try:
        text = "".join(p.get("text", "") for p in data["candidates"][0]["content"]["parts"])
    except (KeyError, IndexError):
        raise LLMError("Gemini returned no answer")
    return json.loads(text)


def unescape(v):
    """Some models return line breaks as a literal backslash-n; turn them back into real ones."""
    if isinstance(v, str):
        return v.replace("\\r\\n", "\n").replace("\\n", "\n").replace("\\t", "  ")
    if isinstance(v, list):
        return [unescape(x) for x in v]
    if isinstance(v, dict):
        return {k: unescape(x) for k, x in v.items()}
    return v


def _built_email(e, name):
    first = (name or "there").split()[0]
    if e.get("paragraphs"):
        return {"subject": e["subject"], "body": compose_email(first, e["paragraphs"])}
    return {"subject": e["subject"], "body": tidy_email(e.get("body", ""), first)}  # model ignored the schema


def score_cv(cv_text, role_hint=None):
    """Call the model, then compute totals/band in code and verify every evidence quote."""
    out = unescape(call_llm(cv_text, role_hint))
    rows, total = [], 0
    for key, label, weight in PARAMS:
        p = out["scores"][key]
        score = max(0, min(3, int(p.get("score", 0))))
        evidence = (p.get("evidence") or "").strip()
        verified = evidence_found(evidence, cv_text) if evidence else None
        rows.append({"key": key, "label": label, "weight": weight, "score": score,
                     "weighted": score * weight, "evidence": evidence, "verified": verified,
                     "reasoning": p.get("reasoning", "")})
        total += score * weight
    return {
        "candidate": out["candidate"],
        "params": rows,
        "total": total,
        "pct": round(100 * total / MAX_WEIGHTED),
        "band": band_for(total),
        "why_ranked": out["why_ranked"],
        "brief": out["brief"],
        "emails": {k: _built_email(out[f"{k}_email"], out["candidate"].get("name")) for k in ("invite", "decline")},
        "unverified": sum(1 for r in rows if r["score"] > 0 and r["verified"] is False),
    }


# ---------- storage ----------

def now():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


class Supabase:
    """Minimal Supabase client over REST: one `candidates` table (id, data jsonb) and a private `resumes` bucket."""

    def __init__(self, url, key):
        self.url = url.rstrip("/")
        self.h = {"apikey": key}
        if key.startswith("eyJ"):  # legacy JWT keys also go in Authorization; new sb_secret_ keys must not
            self.h["Authorization"] = f"Bearer {key}"

    def _req(self, method, path, what, tries=4, **kw):
        """Send a request, retrying Supabase/Cloudflare 5xx and connection errors with backoff."""
        kw.setdefault("timeout", 60)
        kw["headers"] = {**self.h, **kw.get("headers", {})}
        for i in range(tries):
            try:
                r = requests.request(method, f"{self.url}{path}", **kw)
                if r.status_code < 500 or i == tries - 1:
                    break
            except requests.RequestException as e:
                if i == tries - 1:
                    raise RuntimeError(f"Supabase {what} failed: {e}")
            time.sleep(1.5 * 2 ** i)
        if r.status_code >= 300:
            detail = "Supabase is having a temporary problem" if r.status_code >= 500 else r.text[:200]
            raise RuntimeError(f"Supabase {what} failed ({r.status_code}): {detail}")
        return r

    def load_all(self):
        r = self._req("GET", "/rest/v1/candidates?select=id,data", "load")
        return {row["id"]: row["data"] for row in r.json()}

    def upsert(self, c):
        self._req("POST", "/rest/v1/candidates", "save", headers={"Prefer": "resolution=merge-duplicates,return=minimal"},
                  json={"id": c["id"], "data": c, "updated_at": now()})

    def delete(self, cid):
        self._req("DELETE", f"/rest/v1/candidates?id=eq.{cid}", "delete")

    def put_file(self, path, blob, ctype):
        self._req("POST", f"/storage/v1/object/resumes/{path}", "file upload", data=blob,
                  headers={"Content-Type": ctype, "x-upsert": "true"})

    def get_file(self, path):
        return self._req("GET", f"/storage/v1/object/resumes/{path}", "file download").content

    def delete_file(self, path):
        try:
            self._req("DELETE", f"/storage/v1/object/resumes/{path}", "file delete", tries=2)
        except RuntimeError:
            pass


def supabase():
    url, key = env("SUPABASE_URL"), env("SUPABASE_SECRET_KEY")
    return Supabase(url, key) if url and key else None


class Store:
    """Candidates held in memory; persisted to Supabase when configured, else to data/candidates.json."""

    def __init__(self):
        self.lock = threading.Lock()
        UPLOADS.mkdir(parents=True, exist_ok=True)
        self.sb = supabase()
        self.unsynced = set()
        self.items = {}
        if SERVERLESS:
            if not self.sb:
                raise RuntimeError("On Vercel, SUPABASE_URL and SUPABASE_SECRET_KEY must be set")
            return  # loaded per request by refresh()
        if self.sb:
            self.items = self.sb.load_all()
        else:
            self.items = json.loads(DB_FILE.read_text(encoding="utf-8")) if DB_FILE.exists() else {}
        for c in self.items.values():  # jobs interrupted by a restart
            if c["status"] in ("queued", "scoring"):
                c["status"] = "error"
                c["error"] = "Interrupted by a server restart - click Rescore"

    def refresh(self):
        """Serverless: reload from Supabase so every request sees changes made by other instances."""
        items = self.sb.load_all()
        cutoff = time.time() - 360  # a scoring run longer than this was cut off by the function time limit
        for c in items.values():
            if c["status"] in ("queued", "scoring") and c.get("scoring_started", 0) < cutoff:
                c["status"], c["error"] = "error", "Scoring was cut off - click Rescore"
        with self.lock:
            self.items = items

    def _save(self, cid=None):
        if self.sb:
            if cid in self.items:
                self.unsynced.add(cid)
            for k in list(self.unsynced):  # also retries rows a previous Supabase hiccup left unsaved
                if k not in self.items:
                    self.unsynced.discard(k)
                    continue
                try:
                    self.sb.upsert(self.items[k])
                    self.unsynced.discard(k)
                except RuntimeError as e:
                    print(f"warning: will retry saving {k} to Supabase ({e})", flush=True)
                    break
            return
        tmp = DB_FILE.with_suffix(".tmp")
        tmp.write_text(json.dumps(self.items, indent=1, ensure_ascii=False), encoding="utf-8")
        tmp.replace(DB_FILE)

    def get(self, cid):
        with self.lock:
            c = self.items.get(cid)
            return json.loads(json.dumps(c)) if c else None

    def all(self):
        with self.lock:
            return json.loads(json.dumps(list(self.items.values())))

    def update(self, cid, **fields):
        with self.lock:
            self.items[cid].update(fields)
            self._save(cid)
            return json.loads(json.dumps(self.items[cid]))

    def by_hash(self, h):
        with self.lock:
            return next((c["id"] for c in self.items.values() if c["sha"] == h), None)

    def add(self, c):
        with self.lock:
            self.items[c["id"]] = c
            self._save(c["id"])

    def delete(self, cid):
        with self.lock:
            c = self.items.pop(cid, None)
            if self.sb:
                self.sb.delete(cid)
            else:
                self._save()
        if c:
            (UPLOADS / c["stored_as"]).unlink(missing_ok=True)
            if self.sb:
                self.sb.delete_file(c["stored_as"])


def resume_path(c):
    """Local path of the uploaded resume, fetched from Supabase Storage if this machine doesn't have it."""
    p = UPLOADS / c["stored_as"]
    if not p.exists() and store.sb:
        p.write_bytes(store.sb.get_file(c["stored_as"]))
    return p


store = Store()
pool = None if SERVERLESS else ThreadPoolExecutor(max_workers=int(env("SCORING_WORKERS", "4")))


def role_hint_from_filename(name):
    n = name.lower()
    if n.startswith("spm_") or "senior" in n:
        return "Sr PM"
    if n.startswith("pm_"):
        return "PM"
    return None


def ingest(filename, blob):
    """Save an uploaded file and queue it for scoring (its email goes out once scored). Returns (id, is_duplicate)."""
    sha = hashlib.sha256(blob).hexdigest()
    existing = store.by_hash(sha)
    if existing:
        c = store.get(existing)
        # Re-uploading a scored candidate who was never emailed sends their email now; nobody is emailed twice.
        if c["status"] == "done" and not c["sent"] and env("AUTO_SEND", "1") != "0":
            try:
                email_candidate(existing, c["decision"] or "hold")
                store.update(existing, auto_email_error=None)
            except Exception as e:
                store.update(existing, auto_email_error=str(e)[:300])
        return existing, True
    cid = uuid.uuid4().hex[:10]
    safe = re.sub(r"[^A-Za-z0-9._-]+", "_", Path(filename).name)
    stored = f"{cid}_{safe}"
    (UPLOADS / stored).write_bytes(blob)
    file_in_cloud = False
    if store.sb:
        import mimetypes
        try:  # a storage hiccup must not block scoring; the local copy is kept either way
            store.sb.put_file(stored, blob, mimetypes.guess_type(stored)[0] or "application/octet-stream")
            file_in_cloud = True
        except RuntimeError as e:
            print(f"warning: resume file kept locally only ({e})", flush=True)
    store.add({"id": cid, "filename": Path(filename).name, "stored_as": stored, "sha": sha, "file_in_cloud": file_in_cloud,
               "uploaded_at": now(), "status": "queued", "error": None,
               "role_hint": role_hint_from_filename(filename), "result": None,
               "decision": None, "drafts": None, "sent": [], "scoring_started": time.time(),
               # new uploads email the candidate as soon as they are scored (AUTO_SEND=0 turns this off)
               "auto_send": env("AUTO_SEND", "1") != "0"})
    _run(cid)
    return cid, False


def _run(cid):
    """Score in the background locally; on Vercel, score now (the function would be frozen after responding)."""
    if SERVERLESS:
        _score_job(cid)
    else:
        pool.submit(_score_job, cid)


def rescore(cid, role=None):
    c = store.get(cid)
    if not c:
        return None
    if role:
        store.update(cid, role_hint=role)
    store.update(cid, status="queued", error=None, scoring_started=time.time())
    _run(cid)
    return store.get(cid)


def _score_job(cid):
    c = store.get(cid)
    if not c:
        return
    store.update(cid, status="scoring", scoring_started=time.time())
    try:
        text = extract_text(resume_path(c))
        res = score_cv(text, c["role_hint"])
        decision = {"SHORTLIST": "invite", "NOT MOVING FORWARD": "decline"}.get(res["band"], "hold")
        store.update(cid, status="done", result=res, scored_at=now(),
                     decision=decision, drafts={**res["emails"], "hold": hold_draft(res)})
    except Exception as e:  # surface every failure in the dashboard
        # auto_send stays pending: a successful Rescore of an upload that was never emailed still sends it
        store.update(cid, status="error", error=str(e)[:500])
        return
    if store.get(cid).get("auto_send"):
        try:
            email_candidate(cid, decision)
            store.update(cid, auto_send=False, auto_email_error=None)
        except Exception as e:  # scoring succeeded; keep the result and show why the email didn't go
            store.update(cid, auto_send=False, auto_email_error=str(e)[:300])


def hold_draft(res):
    """'Still under review' email for borderline candidates (written here, not by the model, so every borderline gets one)."""
    cand = res["candidate"]
    first = (cand.get("name") or "there").split()[0]
    role = "Senior Product Manager" if cand.get("role_assessed") == "Sr PM" else "Product Manager"
    company = env("COMPANY_NAME", "Kargo")
    return {"subject": f"Your application for {role} at {company}",
            "body": compose_email(first, [
                f"Thank you for applying for the {role} role at {company}.",
                "Your application is still under review. We're meeting other candidates over the next couple of weeks "
                "and will get back to you with a decision as soon as we can.",
                "Thank you for your patience."])}


EMAIL_RE = re.compile(r"[^@\s,;]+@[^@\s,;]+\.[^@\s,;]+")


def email_candidate(cid, kind, to=None, subject=None, body=None, force=False):
    """Send the invite / decline / hold email to a candidate and record it. Returns the updated record."""
    c = store.get(cid)
    if kind not in ("invite", "decline", "hold"):
        raise ValueError("kind must be invite, decline or hold")
    draft = (c.get("drafts") or {}).get(kind) or (hold_draft(c["result"]) if kind == "hold" and c.get("result") else None)
    if not draft and (subject is None or body is None):
        raise ValueError("No email draft for this candidate")
    to = (to if to is not None else (c.get("email_override") or (c.get("result") or {}).get("candidate", {}).get("email") or "")).strip()
    subject = subject if subject is not None else draft["subject"]
    body = body if body is not None else draft["body"]
    if not EMAIL_RE.fullmatch(to):
        raise ValueError("No valid email address for this candidate")
    test = bool(test_recipient())
    if not test and not force and any(not s.get("test") for s in c["sent"]):
        raise PermissionError("An email was already sent to this candidate")
    # same key for a repeated click on the same email; a deliberate re-send gets a new one
    msg_id = send_email(to, subject, body, idempotency_key=f"{cid}-{kind}-{len(c['sent'])}{'-test' if test else ''}")
    drafts = {**(c.get("drafts") or {}), kind: {"subject": subject, "body": body}}
    sent = c["sent"] + [{"kind": kind, "to": to, "at": now(), "via": email_provider(), "id": msg_id,
                         **({"test": True, "delivered_to": test_recipient()} if test else {})}]
    return store.update(cid, sent=sent, drafts=drafts, decision=kind)


# ---------- email ----------

def email_provider():
    if env("RESEND_API_KEY"):
        return "resend"
    if all(env(k) for k in ("SMTP_HOST", "SMTP_USER", "SMTP_PASSWORD")):
        return "smtp"
    return None


def smtp_ready():  # "can send email" - kept under this name for the dashboard config
    return email_provider() is not None


def test_recipient():
    return env("EMAIL_TEST_TO")


def send_email(to, subject, body, idempotency_key=None):
    """Send one email. In test mode (EMAIL_TEST_TO set) it goes to that address instead, labelled with the real recipient."""
    if test_recipient():
        subject = f"[TEST for {to}] {subject}"
        body = (f"--- TEST MODE: this email was meant for {to}. It was sent to you instead. ---\n\n{body}")
        to = test_recipient()
    provider = email_provider()
    if provider == "resend":
        return _send_resend(to, subject, body, idempotency_key)
    if provider is None:
        raise RuntimeError("Email is not set up - add RESEND_API_KEY (or the SMTP_* settings) to .env")
    _send_smtp(to, subject, body)


def _send_resend(to, subject, body, idempotency_key):
    """Send through the Resend API. The idempotency key stops a repeated click from delivering twice."""
    sender = env("RESEND_FROM", "onboarding@resend.dev")
    if "<" not in sender:
        sender = f'{env("SENDER_NAME", "Arjun")} <{sender}>'
    payload = {"from": sender, "to": [to], "subject": subject, "text": body}
    if env("SMTP_BCC"):
        payload["bcc"] = [env("SMTP_BCC")]
    if env("REPLY_TO"):
        payload["reply_to"] = env("REPLY_TO")
    headers = {"Authorization": f'Bearer {env("RESEND_API_KEY")}'}
    if idempotency_key:
        headers["Idempotency-Key"] = idempotency_key
    try:
        r = requests.post("https://api.resend.com/emails", json=payload, headers=headers, timeout=30)
    except requests.RequestException as e:
        raise RuntimeError(f"Could not reach Resend: {e}")
    if r.status_code >= 300:
        try:
            msg = r.json().get("message") or r.text
        except ValueError:
            msg = r.text
        raise RuntimeError(f"Resend refused the email ({r.status_code}): {msg[:250]}")
    return r.json().get("id")


def _send_smtp(to, subject, body):
    msg = EmailMessage()
    from_addr = env("SMTP_FROM", env("SMTP_USER"))
    msg["From"] = f'{env("SENDER_NAME", "Arjun")} <{from_addr}>'
    msg["To"] = to
    msg["Subject"] = subject
    if env("SMTP_BCC"):
        msg["Bcc"] = env("SMTP_BCC")
    msg.set_content(body)
    host, port = env("SMTP_HOST"), int(env("SMTP_PORT", "587"))
    if port == 465:
        with smtplib.SMTP_SSL(host, port, timeout=30) as s:
            s.login(env("SMTP_USER"), env("SMTP_PASSWORD"))
            s.send_message(msg)
    else:
        with smtplib.SMTP(host, port, timeout=30) as s:
            s.starttls()
            s.login(env("SMTP_USER"), env("SMTP_PASSWORD"))
            s.send_message(msg)
